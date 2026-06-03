import io
import csv
import re
import logging
from collections import defaultdict

logger = logging.getLogger(__name__)
from fastapi import APIRouter, Depends, Query, Response, UploadFile, File, HTTPException
from sqlalchemy.orm import Session
from sqlalchemy import desc, asc, or_
from typing import Optional, List
from datetime import date, timedelta
from server.database import get_db
from server.models import Company, StatusHistory, MemoTemplate, ActivityLog, CompanyTag, User, Organization, Plan, EmailSendLog
from server.services.scorer import calculate_score, calculate_digital_maturity
from datetime import datetime as _dt
from server.services.scraper import scrape_company_info
from server.services.categorizer import categorize_company, detect_flags
from server.schemas import company_to_dict
from server.services.cache import cache_invalidate
from server.services.collector import _upsert_company_master
from server.auth import get_current_user, require_phase0_unlock


def _normalize_domain(domain: str) -> str:
    if not domain:
        return ""
    d = domain.lower().strip()
    d = re.sub(r'^(https?://)', '', d)
    d = d.split('/')[0]
    d = re.sub(r'^(www\.)', '', d)
    return d

router = APIRouter(prefix="/api/companies", tags=["companies"])


def _owned_projects(current_user: User, db: Session):
    from server.models import Project
    return [p.id for p in db.query(Project.id).filter(Project.org_id == current_user.org_id).all()]


@router.get("")
def list_companies(
    category: Optional[str] = None,
    status: Optional[str] = None,
    score_rank: Optional[str] = None,
    has_contact: Optional[bool] = None,
    search: Optional[str] = None,
    tag: Optional[str] = None,
    assignee_id: Optional[int] = None,
    follow_up_filter: Optional[str] = None,
    cms_type: Optional[str] = None,
    ec_only: Optional[bool] = None,
    ec_scale: Optional[str] = None,
    sort_by: str = "score_total",
    sort_order: str = "desc",
    page: int = 1,
    per_page: int = 50,
    project_id: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = db.query(Company)
    owned_ids = _owned_projects(current_user, db)
    if project_id:
        if project_id not in owned_ids:
            raise HTTPException(status_code=403, detail="アクセス権限がありません")
        query = query.filter(Company.project_id == project_id)
    else:
        query = query.filter(
            or_(
                Company.project_id.in_(owned_ids),
                Company.org_id == current_user.org_id,
            )
        )

    if category:
        query = query.filter(Company.category_main == category)
    if status:
        query = query.filter(Company.status == status)
    if score_rank:
        query = query.filter(Company.score_rank == score_rank)
    if has_contact is not None:
        if has_contact:
            query = query.filter(Company.contact_url.isnot(None), Company.contact_url != "")
        else:
            query = query.filter((Company.contact_url.is_(None)) | (Company.contact_url == ""))
    if search:
        query = query.filter(
            Company.company_name.ilike(f"%{search}%")
            | Company.website_url.ilike(f"%{search}%")
            | Company.domain.ilike(f"%{search}%")
            | Company.notes.ilike(f"%{search}%")
        )
    if tag:
        tagged_ids = db.query(CompanyTag.company_id).filter(CompanyTag.tag_name == tag).subquery()
        query = query.filter(Company.id.in_(tagged_ids))
    if assignee_id is not None:
        if assignee_id == 0:
            query = query.filter(Company.assignee_id.is_(None))
        else:
            query = query.filter(Company.assignee_id == assignee_id)
    if cms_type:
        if cms_type == "EC_PLATFORMS":
            ec_platforms = ["Shopify", "WooCommerce", "BASE", "MakeShop", "futureshop", "ecbeing",
                            "カラーミー", "EC-CUBE", "STORES", "ロリポップEC", "NEXT ENGINE",
                            "カート365", "Yahoo!ショッピング", "楽天市場", "aishipR", "ショップサーブ"]
            query = query.filter(Company.cms_type.in_(ec_platforms))
        else:
            query = query.filter(Company.cms_type == cms_type)
    if ec_only:
        query = query.filter(Company.ec_flag == True)
    if ec_scale:
        query = query.filter(Company.ec_scale == ec_scale)
    if follow_up_filter:
        today = date.today()
        if follow_up_filter == "overdue":
            query = query.filter(Company.follow_up_date.isnot(None), Company.follow_up_date < today)
        elif follow_up_filter == "today":
            query = query.filter(Company.follow_up_date == today)
        elif follow_up_filter == "week":
            query = query.filter(
                Company.follow_up_date.isnot(None),
                Company.follow_up_date >= today,
                Company.follow_up_date <= today + timedelta(days=7),
            )

    sort_col = getattr(Company, sort_by, Company.score_total)
    if sort_order == "asc":
        query = query.order_by(asc(sort_col))
    else:
        query = query.order_by(desc(sort_col))

    total = query.count()
    companies = query.offset((page - 1) * per_page).limit(per_page).all()

    return {
        "total": total,
        "page": page,
        "per_page": per_page,
        "companies": [company_to_dict(c, db) for c in companies],
    }


@router.post("")
def create_company(
    data: dict,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from server.routes.plans import check_plan_limit
    check_plan_limit(current_user.org_id, "companies", db)

    owned_ids = _owned_projects(current_user, db)

    if data.get("project_id") and data["project_id"] not in owned_ids:
        raise HTTPException(status_code=403, detail="アクセス権限がありません")

    existing = (
        db.query(Company)
        .filter(
            Company.website_url == data.get("website_url"),
            Company.project_id.in_(owned_ids),
        )
        .first()
    )
    if existing:
        return {"error": "この企業URLは既に登録されています", "company": company_to_dict(existing, db)}

    score, rank = calculate_score(data, db=db)
    data["score_total"] = score
    data["score_rank"] = rank

    company = Company(**{k: v for k, v in data.items() if hasattr(Company, k)})
    db.add(company)
    db.commit()
    db.refresh(company)
    cache_invalidate("dashboard")

    try:
        from server.routes.webhooks import fire_event
        fire_event(db, current_user.org_id, "company.created", {
            "company_id": company.id,
            "company_name": company.company_name,
            "domain": company.domain,
            "score_rank": company.score_rank,
            "score_total": company.score_total,
        })
    except Exception:
        pass

    if rank == "A":
        try:
            from server.services.slack_notifier import notify_rank_a_company
            notify_rank_a_company(db, company.company_name or "", domain=company.domain, project_name=None)
        except Exception:
            pass

    domain = _normalize_domain(company.domain or company.website_url or "")
    if domain:
        _upsert_company_master(db, {
            "company_name": company.company_name,
            "website_url": company.website_url,
            "phone": company.phone,
            "email": company.email,
            "prefecture": company.prefecture,
            "city": company.city,
            "category_main": company.category_main,
            "score_total": company.score_total,
            "score_rank": company.score_rank,
        }, domain, source="manual")

    return {"company": company_to_dict(company, db)}


@router.get("/duplicates")
def find_duplicates(
    project_id: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    q = db.query(Company)
    owned_ids = _owned_projects(current_user, db)
    if project_id:
        if project_id not in owned_ids:
            raise HTTPException(status_code=403, detail="アクセス権限がありません")
        q = q.filter(Company.project_id == project_id)
    else:
        q = q.filter(Company.project_id.in_(owned_ids))
    companies = q.all()
    domain_groups = defaultdict(list)
    for c in companies:
        normalized = _normalize_domain(c.domain or c.website_url or "")
        if normalized:
            domain_groups[normalized].append(c)

    duplicate_groups = []
    for norm_domain, group in domain_groups.items():
        if len(group) >= 2:
            duplicate_groups.append({
                "normalized_domain": norm_domain,
                "companies": [company_to_dict(c, db) for c in group],
            })

    return {"duplicate_groups": duplicate_groups, "total_groups": len(duplicate_groups)}


@router.post("/merge")
def merge_companies(
    data: dict,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    main_id = data.get("main_id")
    merge_ids = data.get("merge_ids", [])
    if not main_id or not merge_ids:
        return {"error": "main_id と merge_ids は必須です"}

    owned_ids = _owned_projects(current_user, db)
    main_company = db.query(Company).filter(
        Company.id == main_id, Company.project_id.in_(owned_ids)
    ).first()
    if not main_company:
        return {"error": "メイン企業が見つかりません"}

    merge_companies_list = db.query(Company).filter(
        Company.id.in_(merge_ids), Company.project_id.in_(owned_ids)
    ).all()
    if not merge_companies_list:
        return {"error": "マージ対象の企業が見つかりません"}

    for mc in merge_companies_list:
        if not main_company.phone and mc.phone:
            main_company.phone = mc.phone
        if not main_company.email and mc.email:
            main_company.email = mc.email
        if not main_company.contact_url and mc.contact_url:
            main_company.contact_url = mc.contact_url
        if not main_company.prefecture and mc.prefecture:
            main_company.prefecture = mc.prefecture
        if not main_company.city and mc.city:
            main_company.city = mc.city
        if not main_company.category_main and mc.category_main:
            main_company.category_main = mc.category_main
        if mc.shopify_flag:
            main_company.shopify_flag = True
        if mc.ec_flag:
            main_company.ec_flag = True
        if mc.notes and mc.notes.strip():
            existing_notes = main_company.notes or ""
            if existing_notes:
                main_company.notes = existing_notes + "\n---\n" + mc.notes
            else:
                main_company.notes = mc.notes

        tags = db.query(CompanyTag).filter(CompanyTag.company_id == mc.id).all()
        for tag in tags:
            existing = db.query(CompanyTag).filter(
                CompanyTag.company_id == main_id,
                CompanyTag.tag_name == tag.tag_name,
            ).first()
            if not existing:
                new_tag = CompanyTag(company_id=main_id, tag_name=tag.tag_name)
                db.add(new_tag)

        db.query(CompanyTag).filter(CompanyTag.company_id == mc.id).delete()
        db.query(StatusHistory).filter(StatusHistory.company_id == mc.id).delete()
        db.query(ActivityLog).filter(ActivityLog.company_id == mc.id).delete()
        db.delete(mc)

    company_dict = company_to_dict(main_company, db)
    score, rank = calculate_score(company_dict, db=db)
    main_company.score_total = score
    main_company.score_rank = rank

    db.commit()
    db.refresh(main_company)
    cache_invalidate("dashboard")
    return {"company": company_to_dict(main_company, db), "merged_count": len(merge_companies_list)}


@router.get("/tags/all")
def get_all_tags(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    owned_ids = _owned_projects(current_user, db)
    owned_company_ids = db.query(Company.id).filter(Company.project_id.in_(owned_ids)).subquery()
    tags = (
        db.query(CompanyTag.tag_name)
        .filter(CompanyTag.company_id.in_(owned_company_ids))
        .distinct()
        .order_by(CompanyTag.tag_name)
        .all()
    )
    return {"tags": [t[0] for t in tags]}


@router.get("/pipeline")
def get_pipeline(
    project_id: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    statuses = [
        "未確認", "対象候補", "除外", "アプローチ前",
        "フォーム送信済", "返信あり", "面談化", "代理店化", "失注",
    ]

    owned = _owned_projects(current_user, db)
    if project_id:
        if project_id not in owned:
            raise HTTPException(status_code=403, detail="アクセス権限がありません")
        query = db.query(Company).filter(Company.project_id == project_id)
    else:
        query = db.query(Company).filter(Company.project_id.in_(owned))

    companies = query.order_by(Company.score_total.desc()).all()

    grouped: dict = {s: [] for s in statuses}

    for c in companies:
        card = {
            "id": c.id,
            "company_name": c.company_name or "",
            "domain": c.domain or "",
            "score_total": c.score_total or 0,
            "score_rank": c.score_rank or "D",
            "status": c.status or "未確認",
            "follow_up_date": c.follow_up_date.isoformat() if c.follow_up_date else None,
            "assignee_id": c.assignee_id,
            "ec_flag": c.ec_flag,
            "shopify_flag": c.shopify_flag,
            "cms_type": c.cms_type,
            "category_main": c.category_main or "",
            "contact_name": c.contact_name or "",
        }
        target = c.status if c.status in grouped else "未確認"
        grouped[target].append(card)

    counts = {s: len(grouped[s]) for s in statuses}
    return {
        "columns": [{"status": s, "companies": grouped[s]} for s in statuses],
        "counts": counts,
        "total": len(companies),
    }


@router.get("/export.xlsx")
def export_companies_xlsx_v2(
    category: Optional[str] = None,
    status: Optional[str] = None,
    score_rank: Optional[str] = None,
    search: Optional[str] = None,
    project_id: Optional[int] = None,
    cms_type: Optional[str] = None,
    ec_only: Optional[bool] = None,
    ec_scale: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from io import BytesIO
    from fastapi.responses import StreamingResponse
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    query = db.query(Company)
    owned_ids = _owned_projects(current_user, db)
    if project_id:
        if project_id not in owned_ids:
            raise HTTPException(status_code=403, detail="アクセス権限がありません")
        query = query.filter(Company.project_id == project_id)
    else:
        query = query.filter(
            or_(
                Company.project_id.in_(owned_ids),
                Company.org_id == current_user.org_id,
            )
        )

    if category:
        query = query.filter(Company.category_main == category)
    if status:
        query = query.filter(Company.status == status)
    if score_rank:
        query = query.filter(Company.score_rank == score_rank)
    if search:
        query = query.filter(
            Company.company_name.ilike(f"%{search}%")
            | Company.website_url.ilike(f"%{search}%")
            | Company.domain.ilike(f"%{search}%")
        )
    if cms_type:
        if cms_type == "EC_PLATFORMS":
            ec_platforms = ["Shopify", "WooCommerce", "BASE", "MakeShop", "futureshop", "ecbeing",
                            "カラーミー", "EC-CUBE", "STORES", "ロリポップEC", "NEXT ENGINE",
                            "カート365", "Yahoo!ショッピング", "楽天市場", "aishipR", "ショップサーブ"]
            query = query.filter(Company.cms_type.in_(ec_platforms))
        else:
            query = query.filter(Company.cms_type == cms_type)
    if ec_only:
        query = query.filter(Company.ec_flag == True)
    if ec_scale:
        query = query.filter(Company.ec_scale == ec_scale)

    companies = query.order_by(desc(Company.score_total)).limit(5000).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "企業リスト"

    headers = ["会社名", "ドメイン", "WebサイトURL", "問い合わせURL", "電話番号", "メールアドレス",
               "都道府県", "市区町村", "カテゴリ", "CMS種別", "ECサイト", "ステータス", "スコア", "ランク",
               "担当者", "フォローアップ日", "メモ", "登録日"]
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for c in companies:
        ws.append([
            c.company_name or "",
            c.domain or "",
            c.website_url or "",
            c.contact_url or "",
            c.phone or "",
            c.email or "",
            c.prefecture or "",
            c.city or "",
            c.category_main or "",
            c.cms_type or "",
            "○" if c.ec_flag else "",
            c.status or "",
            c.score_total or 0,
            c.score_rank or "",
            "",
            c.follow_up_date.strftime("%Y-%m-%d") if c.follow_up_date else "",
            c.notes or "",
            c.created_at.strftime("%Y-%m-%d") if c.created_at else "",
        ])

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=companies_export.xlsx"},
    )


@router.get("/ids")
def get_company_ids(
    category: Optional[str] = None,
    status: Optional[str] = None,
    score_rank: Optional[str] = None,
    has_contact: Optional[bool] = None,
    search: Optional[str] = None,
    tag: Optional[str] = None,
    assignee_id: Optional[int] = None,
    follow_up_filter: Optional[str] = None,
    cms_type: Optional[str] = None,
    ec_only: Optional[bool] = None,
    ec_scale: Optional[str] = None,
    project_id: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = db.query(Company.id, Company.email)
    owned_ids = _owned_projects(current_user, db)
    if project_id:
        if project_id not in owned_ids:
            raise HTTPException(status_code=403, detail="アクセス権限がありません")
        query = query.filter(Company.project_id == project_id)
    else:
        query = query.filter(
            or_(
                Company.project_id.in_(owned_ids),
                Company.org_id == current_user.org_id,
            )
        )
    if category:
        query = query.filter(Company.category_main == category)
    if status:
        query = query.filter(Company.status == status)
    if score_rank:
        query = query.filter(Company.score_rank == score_rank)
    if has_contact is not None:
        if has_contact:
            query = query.filter(Company.contact_url.isnot(None), Company.contact_url != "")
        else:
            query = query.filter((Company.contact_url.is_(None)) | (Company.contact_url == ""))
    if search:
        query = query.filter(
            Company.company_name.ilike(f"%{search}%")
            | Company.website_url.ilike(f"%{search}%")
            | Company.domain.ilike(f"%{search}%")
            | Company.notes.ilike(f"%{search}%")
        )
    if tag:
        tagged_ids = db.query(CompanyTag.company_id).filter(CompanyTag.tag_name == tag).subquery()
        query = query.filter(Company.id.in_(tagged_ids))
    if assignee_id is not None:
        if assignee_id == 0:
            query = query.filter(Company.assignee_id.is_(None))
        else:
            query = query.filter(Company.assignee_id == assignee_id)
    if cms_type:
        if cms_type == "EC_PLATFORMS":
            ec_platforms = ["Shopify", "WooCommerce", "BASE", "MakeShop", "futureshop", "ecbeing",
                            "カラーミー", "EC-CUBE", "STORES", "ロリポップEC", "NEXT ENGINE",
                            "カート365", "Yahoo!ショッピング", "楽天市場", "aishipR", "ショップサーブ"]
            query = query.filter(Company.cms_type.in_(ec_platforms))
        else:
            query = query.filter(Company.cms_type == cms_type)
    if ec_only:
        query = query.filter(Company.ec_flag == True)
    if ec_scale:
        query = query.filter(Company.ec_scale == ec_scale)
    if follow_up_filter:
        today = date.today()
        if follow_up_filter == "overdue":
            query = query.filter(Company.follow_up_date.isnot(None), Company.follow_up_date < today)
        elif follow_up_filter == "today":
            query = query.filter(Company.follow_up_date == today)
        elif follow_up_filter == "week":
            query = query.filter(
                Company.follow_up_date.isnot(None),
                Company.follow_up_date >= today,
                Company.follow_up_date <= today + timedelta(days=7),
            )
    MAX_BULK = 500
    rows = query.limit(MAX_BULK).all()
    ids = [r[0] for r in rows]
    with_email = sum(1 for r in rows if r[1])
    return {
        "ids": ids,
        "total": len(ids),
        "with_email": with_email,
        "capped": len(rows) == MAX_BULK,
    }


@router.get("/for-sales-ai")
def get_companies_for_sales_ai(
    project_id: Optional[int] = None,
    show_all: Optional[bool] = False,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """SalesAI ターゲット選択用スリムエンドポイント。必要9フィールドのみ返し、N+1クエリを排除。"""
    owned_ids = _owned_projects(current_user, db)
    query = db.query(
        Company.id,
        Company.company_name,
        Company.score_total,
        Company.score_rank,
        Company.ec_flag,
        Company.cms_type,
        Company.prefecture,
        Company.category_main,
        Company.email,
    )
    if project_id and not show_all:
        if project_id not in owned_ids:
            raise HTTPException(status_code=403, detail="アクセス権限がありません")
        query = query.filter(Company.project_id == project_id)
    elif show_all:
        query = query.filter(
            or_(
                Company.project_id.in_(owned_ids),
                Company.org_id == current_user.org_id,
            )
        )
    else:
        query = query.filter(
            or_(
                Company.project_id.in_(owned_ids),
                Company.org_id == current_user.org_id,
            )
        )

    rows = query.order_by(desc(Company.score_total)).all()
    companies = [
        {
            "id": r[0],
            "company_name": r[1],
            "score_total": r[2] or 0,
            "score_rank": r[3] or "D",
            "ec_flag": bool(r[4]),
            "cms_type": r[5],
            "prefecture": r[6],
            "category_main": r[7],
            "email": r[8],
        }
        for r in rows
    ]
    categories = sorted(set(c["category_main"] for c in companies if c["category_main"]))
    return {"companies": companies, "total": len(companies), "categories": categories}


@router.post("/bulk-scan-forms")
def bulk_scan_forms(
    data: dict,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """選択した企業のお問い合わせフォームURLを一括スキャンしてDBに保存する。"""
    import requests as _req
    from server.services.form_sender import _find_contact_url
    company_ids = data.get("company_ids", [])
    if not company_ids:
        return {"scanned": 0, "found": 0, "not_found": 0}

    owned_ids = _owned_projects(current_user, db)
    companies = db.query(Company).filter(
        Company.id.in_(company_ids),
        Company.project_id.in_(owned_ids),
    ).all()

    found_count = 0
    not_found_count = 0
    session = _req.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124.0.0.0 Safari/537.36",
        "Accept-Language": "ja,en-US;q=0.9,en;q=0.8",
    })

    for company in companies:
        if not company.website_url:
            not_found_count += 1
            continue
        try:
            url = _find_contact_url(company.website_url, "", session)
            if url:
                company.contact_url = url
                found_count += 1
            else:
                not_found_count += 1
        except Exception:
            not_found_count += 1

    db.commit()
    return {
        "scanned": len(companies),
        "found": found_count,
        "not_found": not_found_count,
    }


@router.get("/{company_id}")
def get_company(
    company_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    project_ids = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id,
        Company.project_id.in_(project_ids),
    ).first()
    if not company:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="企業が見つかりません")
    return {"company": company_to_dict(company, db)}


@router.put("/{company_id}")
def update_company(
    company_id: int,
    data: dict,
    current_user: User = Depends(require_phase0_unlock),
    db: Session = Depends(get_db),
):
    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        return {"error": "企業が見つかりません"}

    old_status = company.status

    for key, value in data.items():
        if hasattr(company, key) and key not in ("id", "created_at"):
            setattr(company, key, value)

    new_status = company.status
    if old_status != new_status:
        history = StatusHistory(
            company_id=company.id,
            old_status=old_status,
            new_status=new_status,
            user_id=current_user.id,
            user_name=getattr(current_user, "display_name", None) or getattr(current_user, "email", None),
            note=data.get("status_note"),
        )
        db.add(history)

    company_dict = company_to_dict(company, db)
    score, rank = calculate_score(company_dict, db=db)
    company.score_total = score
    company.score_rank = rank
    company.digital_maturity_score = calculate_digital_maturity(company_dict)
    company.score_updated_at = _dt.utcnow()

    db.commit()
    db.refresh(company)
    cache_invalidate("dashboard")
    return {"company": company_to_dict(company, db)}


@router.delete("/{company_id}")
def delete_company(
    company_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    owned_ids = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id, Company.project_id.in_(owned_ids)
    ).first()
    if not company:
        return {"error": "企業が見つかりません"}
    db.delete(company)
    db.commit()
    cache_invalidate("dashboard")
    return {"message": "削除しました"}


@router.put("/bulk-status")
def bulk_update_status(
    data: dict,
    current_user: User = Depends(require_phase0_unlock),
    db: Session = Depends(get_db),
):
    company_ids = data.get("company_ids", [])
    new_status = data.get("new_status", "")
    if not company_ids or not new_status:
        return {"error": "company_ids と new_status は必須です"}

    companies = db.query(Company).filter(Company.id.in_(company_ids)).all()
    updated = 0
    _user_name = getattr(current_user, "display_name", None) or getattr(current_user, "email", None)
    for company in companies:
        old_status = company.status
        if old_status != new_status:
            company.status = new_status
            history = StatusHistory(
                company_id=company.id,
                old_status=old_status,
                new_status=new_status,
                user_id=current_user.id,
                user_name=_user_name,
            )
            db.add(history)
            updated += 1

    db.commit()
    cache_invalidate("dashboard")
    return {"updated": updated, "total": len(company_ids)}


@router.patch("/{company_id}/status")
def patch_status(
    company_id: int,
    data: dict,
    current_user: User = Depends(require_phase0_unlock),
    db: Session = Depends(get_db),
):
    new_status = data.get("status", "")
    if not new_status:
        raise HTTPException(status_code=400, detail="status は必須です")

    owned = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id,
        Company.project_id.in_(owned),
    ).first()
    if not company:
        raise HTTPException(status_code=404, detail="企業が見つかりません")

    old_status = company.status
    if old_status != new_status:
        company.status = new_status
        db.add(StatusHistory(
            company_id=company.id,
            old_status=old_status,
            new_status=new_status,
            user_id=current_user.id,
            user_name=getattr(current_user, "display_name", None) or getattr(current_user, "email", None),
            note=data.get("note"),
        ))
        db.commit()
        cache_invalidate("dashboard")
        try:
            from server.routes.webhooks import fire_event
            fire_event(db, current_user.org_id, "company.stage_changed", {
                "company_id": company.id,
                "company_name": company.company_name,
                "old_status": old_status,
                "new_status": new_status,
            })
        except Exception:
            pass
    return {"id": company.id, "status": company.status}


@router.get("/{company_id}/history")
def get_status_history(
    company_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    owned_ids = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id, Company.project_id.in_(owned_ids)
    ).first()
    if not company:
        raise HTTPException(status_code=404, detail="企業が見つかりません")
    history = (
        db.query(StatusHistory)
        .filter(StatusHistory.company_id == company_id)
        .order_by(desc(StatusHistory.changed_at))
        .all()
    )
    return {
        "history": [
            {
                "id": h.id,
                "old_status": h.old_status,
                "new_status": h.new_status,
                "user_id": h.user_id,
                "user_name": h.user_name,
                "note": h.note,
                "changed_at": h.changed_at.isoformat() if h.changed_at else None,
            }
            for h in history
        ]
    }


@router.get("/{company_id}/activities")
def get_activities(
    company_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    owned_ids = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id, Company.project_id.in_(owned_ids)
    ).first()
    if not company:
        raise HTTPException(status_code=404, detail="企業が見つかりません")
    activities = (
        db.query(ActivityLog)
        .filter(ActivityLog.company_id == company_id)
        .order_by(desc(ActivityLog.created_at))
        .all()
    )
    return {
        "activities": [
            {
                "id": a.id,
                "company_id": a.company_id,
                "action_type": a.action_type,
                "description": a.description,
                "created_at": a.created_at.isoformat() if a.created_at else None,
            }
            for a in activities
        ]
    }


@router.post("/{company_id}/activities")
def create_activity(
    company_id: int,
    data: dict,
    current_user: User = Depends(require_phase0_unlock),
    db: Session = Depends(get_db),
):
    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        return {"error": "企業が見つかりません"}

    activity = ActivityLog(
        company_id=company_id,
        action_type=data.get("action_type", "その他"),
        description=data.get("description", ""),
    )
    db.add(activity)
    db.commit()
    db.refresh(activity)
    return {
        "activity": {
            "id": activity.id,
            "company_id": activity.company_id,
            "action_type": activity.action_type,
            "description": activity.description,
            "created_at": activity.created_at.isoformat() if activity.created_at else None,
        }
    }


@router.get("/csv")
def export_csv(
    category: Optional[str] = None,
    status: Optional[str] = None,
    score_rank: Optional[str] = None,
    has_contact: Optional[bool] = None,
    project_id: Optional[int] = None,
    cms_type: Optional[str] = None,
    ec_only: Optional[bool] = None,
    ec_scale: Optional[str] = None,
    current_user: User = Depends(require_phase0_unlock),
    db: Session = Depends(get_db),
):
    from fastapi import HTTPException as _HTTPException
    from datetime import datetime as _dt

    org = db.query(Organization).filter(Organization.id == current_user.org_id).first()
    csv_limit = None
    if org and org.plan_id:
        plan = db.query(Plan).filter(Plan.id == org.plan_id).first()
        if plan:
            raw = getattr(plan, "max_csv_export", None)
            if raw is not None:
                if raw == 0:
                    raise _HTTPException(
                        status_code=402,
                        detail="CSVエクスポートはスターター以上のプランでご利用いただけます。"
                    )
                csv_limit = raw

    query = db.query(Company)
    owned_ids = _owned_projects(current_user, db)
    if project_id:
        if project_id not in owned_ids:
            raise HTTPException(status_code=403, detail="アクセス権限がありません")
        query = query.filter(Company.project_id == project_id)
    else:
        query = query.filter(
            or_(
                Company.project_id.in_(owned_ids),
                Company.org_id == current_user.org_id,
            )
        )
    if category:
        query = query.filter(Company.category_main == category)
    if status:
        query = query.filter(Company.status == status)
    if score_rank:
        query = query.filter(Company.score_rank == score_rank)
    if has_contact:
        query = query.filter(Company.contact_url.isnot(None), Company.contact_url != "")
    if cms_type:
        if cms_type == "EC_PLATFORMS":
            ec_platforms = ["Shopify", "WooCommerce", "BASE", "MakeShop", "futureshop", "ecbeing",
                            "カラーミー", "EC-CUBE", "STORES", "ロリポップEC", "NEXT ENGINE",
                            "カート365", "Yahoo!ショッピング", "楽天市場", "aishipR", "ショップサーブ"]
            query = query.filter(Company.cms_type.in_(ec_platforms))
        else:
            query = query.filter(Company.cms_type == cms_type)
    if ec_only:
        query = query.filter(Company.ec_flag == True)
    if ec_scale:
        query = query.filter(Company.ec_scale == ec_scale)

    query = query.order_by(desc(Company.score_total))
    if csv_limit is not None:
        query = query.limit(csv_limit)
    companies = query.all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "会社名", "URL", "問い合わせURL", "所在地", "電話番号",
        "メール", "カテゴリ", "CMS種別", "ECサイト", "Shopify対応",
        "スコア", "ランク", "ステータス", "メモ",
    ])

    for c in companies:
        location = f"{c.prefecture or ''}{c.city or ''}"
        writer.writerow([
            c.company_name or "",
            c.website_url or "",
            c.contact_url or "",
            location,
            c.phone or "",
            c.email or "",
            c.category_main or "",
            c.cms_type or "",
            "○" if c.ec_flag else "",
            "○" if c.shopify_flag else "",
            c.score_total,
            c.score_rank or "",
            c.status or "",
            c.notes or "",
        ])

    csv_content = output.getvalue()
    output.close()

    today = _dt.utcnow().strftime("%Y%m%d")
    headers = {
        "Content-Disposition": f"attachment; filename=companies_export_{today}.csv",
        "X-Export-Count": str(len(companies)),
    }
    if csv_limit is not None:
        headers["X-Export-Limit"] = str(csv_limit)

    return Response(
        content=csv_content.encode("utf-8-sig"),
        media_type="text/csv",
        headers=headers,
    )


@router.get("/csv/template")
def download_csv_template(
    current_user: User = Depends(get_current_user),
):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["name", "website_url", "email", "phone", "prefecture", "city", "memo", "status", "category_main", "rank"])
    writer.writerow(["株式会社サンプル", "https://example.com", "info@example.com", "03-1234-5678", "東京都", "渋谷区", "サンプルメモ", "未確認", "EC制作", "A"])
    content = output.getvalue()
    output.close()
    return Response(
        content=content.encode("utf-8-sig"),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=companies_import_template.csv"},
    )


@router.post("/import-csv")
async def import_csv(
    file: UploadFile = File(...),
    project_id: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if not project_id:
        return {"error": "project_id は必須です"}

    project_ids = _owned_projects(current_user, db)
    if project_id not in project_ids:
        return {"error": "プロジェクトへのアクセス権限がありません"}

    content = await file.read()
    try:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("shift_jis", errors="replace")
    except Exception:
        return {"error": "ファイルの読み込みに失敗しました"}

    reader = csv.DictReader(io.StringIO(text))
    added = 0
    skipped = 0
    errors = []
    master_db_entries = []

    COLUMN_MAP = {
        "name": "company_name",
        "会社名": "company_name",
        "website_url": "website_url",
        "URL": "website_url",
        "url": "website_url",
        "email": "email",
        "メール": "email",
        "phone": "phone",
        "電話番号": "phone",
        "prefecture": "prefecture",
        "都道府県": "prefecture",
        "city": "city",
        "市区町村": "city",
        "memo": "notes",
        "メモ": "notes",
        "status": "status",
        "ステータス": "status",
        "category_main": "category_main",
        "カテゴリ": "category_main",
        "rank": "score_rank",
        "ランク": "score_rank",
    }

    VALID_STATUSES = ["未確認", "確認済み", "コンタクト済み", "返信あり", "商談中", "提案済み", "契約交渉中", "代理店化", "不採用"]
    VALID_RANKS = ["A", "B", "C", "D"]

    for i, row in enumerate(reader, start=2):
        try:
            mapped = {}
            for col, val in row.items():
                field = COLUMN_MAP.get(col.strip(), col.strip())
                mapped[field] = val.strip() if val else ""

            url = mapped.get("website_url", "")
            if not url:
                skipped += 1
                continue

            existing = db.query(Company).filter(
                Company.website_url == url,
                Company.project_id == project_id,
            ).first()
            if existing:
                skipped += 1
                continue

            status_val = mapped.get("status", "未確認")
            if status_val not in VALID_STATUSES:
                status_val = "未確認"

            rank_val = mapped.get("score_rank", "").upper()
            if rank_val not in VALID_RANKS:
                rank_val = "D"

            domain_match = re.sub(r'^(https?://)', '', url.lower().strip()).split('/')[0]
            domain_match = re.sub(r'^(www\.)', '', domain_match)

            company = Company(
                project_id=project_id,
                company_name=mapped.get("company_name", ""),
                website_url=url,
                domain=domain_match,
                email=mapped.get("email", ""),
                phone=mapped.get("phone", ""),
                prefecture=mapped.get("prefecture", ""),
                city=mapped.get("city", ""),
                notes=mapped.get("notes", ""),
                status=status_val,
                category_main=mapped.get("category_main", ""),
                score_rank=rank_val,
            )
            db.add(company)
            db.flush()
            added += 1
            master_db_entries.append({
                "domain": domain_match,
                "data": {
                    "company_name": mapped.get("company_name", ""),
                    "website_url": url,
                    "email": mapped.get("email", ""),
                    "phone": mapped.get("phone", ""),
                    "prefecture": mapped.get("prefecture", ""),
                    "city": mapped.get("city", ""),
                    "category_main": mapped.get("category_main", ""),
                    "score_rank": rank_val,
                },
            })
        except Exception as e:
            errors.append(f"行{i}: {str(e)}")

    db.commit()
    cache_invalidate("dashboard")

    for entry in master_db_entries:
        if entry["domain"]:
            _upsert_company_master(db, entry["data"], entry["domain"], source="csv_import")

    return {
        "added": added,
        "skipped": skipped,
        "errors": errors,
        "message": f"{added}件を追加しました（スキップ: {skipped}件）",
    }


@router.get("/{company_id}/tags")
def get_company_tags(
    company_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    owned_ids = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id, Company.project_id.in_(owned_ids)
    ).first()
    if not company:
        raise HTTPException(status_code=404, detail="企業が見つかりません")
    tags = db.query(CompanyTag).filter(CompanyTag.company_id == company_id).all()
    return {"tags": [{"id": t.id, "tag_name": t.tag_name, "created_at": t.created_at.isoformat() if t.created_at else None} for t in tags]}


@router.post("/{company_id}/tags")
def add_company_tag(
    company_id: int,
    data: dict,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    owned_ids = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id, Company.project_id.in_(owned_ids)
    ).first()
    if not company:
        return {"error": "企業が見つかりません"}

    tag_name = data.get("tag_name", "").strip()
    if not tag_name:
        return {"error": "タグ名は必須です"}

    existing = db.query(CompanyTag).filter(
        CompanyTag.company_id == company_id,
        CompanyTag.tag_name == tag_name
    ).first()
    if existing:
        return {"error": "このタグは既に追加されています"}

    tag = CompanyTag(company_id=company_id, tag_name=tag_name)
    db.add(tag)
    db.commit()
    db.refresh(tag)
    return {"tag": {"id": tag.id, "tag_name": tag.tag_name, "created_at": tag.created_at.isoformat() if tag.created_at else None}}


@router.delete("/{company_id}/tags/{tag_name}")
def delete_company_tag(
    company_id: int,
    tag_name: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    owned_ids = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id, Company.project_id.in_(owned_ids)
    ).first()
    if not company:
        return {"error": "企業が見つかりません"}
    tag = db.query(CompanyTag).filter(
        CompanyTag.company_id == company_id,
        CompanyTag.tag_name == tag_name
    ).first()
    if not tag:
        return {"error": "タグが見つかりません"}
    db.delete(tag)
    db.commit()
    return {"message": "タグを削除しました"}


@router.post("/{company_id}/rescrape")
def rescrape_company(
    company_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    owned_ids = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id, Company.project_id.in_(owned_ids)
    ).first()
    if not company:
        return {"error": "企業が見つかりません"}

    url = company.website_url
    if not url:
        return {"error": "WebサイトURLが設定されていません"}

    scraped = scrape_company_info(url)
    if scraped.get("error"):
        return {"error": f"スクレイピングに失敗しました: {scraped['error']}"}

    full_text = scraped.get("full_text", "")
    category_main, category_sub = categorize_company(full_text)
    cms_type = scraped.get("cms_type") or None
    flags = detect_flags(full_text, cms_type=cms_type)

    saved_adjustment = company.score_adjustment or 0

    if scraped.get("company_name"):
        company.company_name = scraped["company_name"]
    if scraped.get("domain"):
        company.domain = scraped["domain"]
    if scraped.get("contact_url"):
        company.contact_url = scraped["contact_url"]
    if scraped.get("phone"):
        company.phone = scraped["phone"]
    if scraped.get("email"):
        company.email = scraped["email"]
    if scraped.get("prefecture"):
        company.prefecture = scraped["prefecture"]
    if scraped.get("city"):
        company.city = scraped["city"]

    company.category_main = category_main
    company.category_sub = category_sub
    for flag_key, flag_val in flags.items():
        setattr(company, flag_key, flag_val)

    company.score_adjustment = saved_adjustment
    company_dict = company_to_dict(company, db)
    score, rank = calculate_score(company_dict, db=db)
    company.score_total = score
    company.score_rank = rank

    db.commit()
    db.refresh(company)
    return {"company": company_to_dict(company, db)}


@router.post("/{company_id}/scan-form")
def scan_company_form(
    company_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """単体企業のお問い合わせフォームURLをスキャンしてDBに保存する。"""
    import requests as _req
    from server.services.form_sender import _find_contact_url
    owned_ids = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id, Company.project_id.in_(owned_ids)
    ).first()
    if not company:
        raise HTTPException(status_code=404, detail="企業が見つかりません")
    if not company.website_url:
        return {"success": False, "message": "WebサイトURLが未設定です", "contact_url": None}

    session = _req.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124.0.0.0 Safari/537.36",
        "Accept-Language": "ja,en-US;q=0.9,en;q=0.8",
    })
    try:
        url = _find_contact_url(company.website_url, "", session)
        if url:
            company.contact_url = url
            db.commit()
            db.refresh(company)
            return {"success": True, "message": f"フォームURLを検出しました", "contact_url": url}
        else:
            return {"success": False, "message": "フォームURLが見つかりませんでした", "contact_url": None}
    except Exception as e:
        return {"success": False, "message": f"スキャンエラー: {str(e)[:60]}", "contact_url": None}


@router.post("/{company_id}/ai-analyze")
def ai_analyze_company(
    company_id: int,
    current_user: User = Depends(require_phase0_unlock),
    db: Session = Depends(get_db),
):
    from server.services.ai_analyzer import analyze_company, get_openai_key
    from server.services.scraper import scrape_company_info
    from server.routes.plans import check_plan_limit

    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        return {"error": "企業が見つかりません"}

    api_key = get_openai_key(db, current_user.org_id)
    if not api_key:
        return {"error": "OpenAI APIキーが設定されていません。設定画面からAPIキーを登録してください。"}

    check_plan_limit(current_user.org_id, "ai_analyses", db)

    full_text = ""
    if company.website_url:
        scraped = scrape_company_info(company.website_url)
        full_text = scraped.get("full_text", "")

    result = analyze_company(
        company_name=company.company_name or company.domain or "",
        url=company.website_url or "",
        full_text=full_text,
        api_key=api_key,
        db=db,
        org_id=current_user.org_id,
        user_id=current_user.id,
    )

    if not result.get("success"):
        return {"error": f"AI分析に失敗しました: {result.get('error', '不明なエラー')}"}

    company.ai_summary = result["summary"]
    db.commit()
    db.refresh(company)
    return {"company": company_to_dict(company, db), "summary": result["summary"]}


@router.post("/{company_id}/generate-email")
def generate_company_email(
    company_id: int,
    data: dict,
    current_user: User = Depends(require_phase0_unlock),
    db: Session = Depends(get_db),
):
    from server.services.ai_analyzer import generate_outreach_email, get_openai_key

    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        return {"error": "企業が見つかりません"}

    api_key = get_openai_key(db, current_user.org_id)
    if not api_key:
        return {"error": "OpenAI APIキーが設定されていません。設定画面からAPIキーを登録してください。"}

    tone = data.get("tone", "formal")
    custom_note = data.get("custom_note", "")

    result = generate_outreach_email(
        company_name=company.company_name or company.domain or "",
        url=company.website_url or "",
        category=company.category_main or "",
        ai_summary=company.ai_summary,
        tone=tone,
        custom_note=custom_note,
        api_key=api_key,
        db=db,
        org_id=current_user.org_id,
        user_id=current_user.id,
    )

    if not result.get("success"):
        return {"error": f"メール生成に失敗しました: {result.get('error', '不明なエラー')}"}

    return {"email": result["email"]}


@router.post("/{company_id}/send-email")
def send_company_email(
    company_id: int,
    data: dict,
    current_user: User = Depends(require_phase0_unlock),
    db: Session = Depends(get_db),
):
    from server.services.mailer import get_smtp_settings, send_email

    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        return {"error": "企業が見つかりません"}

    subject = data.get("subject", "").strip()
    body = data.get("body", "").strip()
    to_email = data.get("to_email", "").strip() or company.email or ""
    template_id = data.get("template_id")

    if not to_email:
        return {"error": "送信先メールアドレスが設定されていません"}
    if not subject:
        return {"error": "件名を入力してください"}
    if not body:
        return {"error": "本文を入力してください"}

    org_id = current_user.org_id

    from server.routes.plans import check_smtp_allowed
    check_smtp_allowed(org_id, db)

    smtp_settings = get_smtp_settings(db, org_id)

    html_body = body.replace("\n", "<br>")
    success, message = send_email(
        to=to_email,
        subject=subject,
        html_body=html_body,
        smtp_settings=smtp_settings,
        text_body=body,
    )

    log = EmailSendLog(
        company_id=company_id,
        org_id=org_id,
        sent_by_id=current_user.id,
        template_id=template_id,
        subject=subject,
        body=body,
        to_email=to_email,
        status="sent" if success else "failed",
        error_message=None if success else message,
    )
    db.add(log)

    if success:
        activity = ActivityLog(
            company_id=company_id,
            action_type="メール送信",
            description=f"件名: {subject} → {to_email}",
        )
        db.add(activity)

    db.commit()

    if success:
        return {"success": True, "message": message}
    return {"error": message}


@router.get("/{company_id}/email-logs")
def get_email_logs(
    company_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    owned_ids = _owned_projects(current_user, db)
    company = db.query(Company).filter(
        Company.id == company_id, Company.project_id.in_(owned_ids)
    ).first()
    if not company:
        raise HTTPException(status_code=404, detail="企業が見つかりません")
    logs = (
        db.query(EmailSendLog)
        .filter(EmailSendLog.company_id == company_id)
        .order_by(desc(EmailSendLog.sent_at))
        .limit(50)
        .all()
    )
    result = []
    for log in logs:
        sender = None
        if log.sent_by_id:
            u = db.query(User).filter(User.id == log.sent_by_id).first()
            if u:
                sender = u.display_name or u.email
        result.append({
            "id": log.id,
            "subject": log.subject,
            "to_email": log.to_email,
            "status": log.status,
            "error_message": log.error_message,
            "sent_by": sender,
            "sent_at": log.sent_at.isoformat() if log.sent_at else None,
        })
    return {"logs": result}


@router.post("/move-project")
def move_project(
    data: dict,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    company_ids = data.get("company_ids", [])
    target_project_id = data.get("target_project_id")
    if not company_ids or not target_project_id:
        return {"error": "企業IDと移動先プロジェクトIDを指定してください"}

    owned_ids = _owned_projects(current_user, db)
    if target_project_id not in owned_ids:
        raise HTTPException(status_code=403, detail="アクセス権限がありません")

    existing_domains = set(
        c.domain for c in db.query(Company.domain).filter(Company.project_id == target_project_id).all()
    )

    moved = 0
    skipped = 0
    for cid in company_ids:
        company = db.query(Company).filter(
            Company.id == cid, Company.project_id.in_(owned_ids)
        ).first()
        if not company:
            continue
        if company.domain in existing_domains:
            skipped += 1
            continue
        company.project_id = target_project_id
        existing_domains.add(company.domain)
        moved += 1

    db.commit()
    return {"moved": moved, "skipped": skipped}


@router.post("/bulk-rescore")
def bulk_rescore_companies(
    data: dict,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from server.services.scorer import calculate_score

    project_id = data.get("project_id")
    q = db.query(Company)
    if project_id:
        q = q.filter(Company.project_id == project_id)
    companies_list = q.all()

    updated = 0
    for company in companies_list:
        try:
            flags = {
                "shopify_flag": company.shopify_flag or False,
                "ec_flag": company.ec_flag or False,
                "amazon_flag": company.amazon_flag or False,
                "rakuten_flag": company.rakuten_flag or False,
                "base_flag": getattr(company, "base_flag", False) or False,
                "makeshop_flag": getattr(company, "makeshop_flag", False) or False,
                "futureshop_flag": getattr(company, "futureshop_flag", False) or False,
                "stores_flag": getattr(company, "stores_flag", False) or False,
                "consulting_flag": company.consulting_flag or False,
                "operation_flag": company.operation_flag or False,
                "production_flag": company.production_flag or False,
            }
            score_total, score_rank = calculate_score(
                flags,
                ec_score=company.ec_score or 0,
                adjustment=company.score_adjustment or 0,
                scoring_rules=None,
            )
            company.score_total = score_total
            company.score_rank = score_rank
            updated += 1
        except Exception as e:
            logger.warning(f"rescore error company {company.id}: {e}")
            continue

    db.commit()
    return {"updated": updated, "total": len(companies_list), "message": f"{updated}件のスコアを再計算しました"}
